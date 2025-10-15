import oracledb as cx_Oracle
import os
import logging
import pandas as pd
from dotenv import load_dotenv
from datetime import datetime, timedelta
import requests
import json
import base64
import re
import io
import csv
import openai
import numpy as np
from typing import Tuple, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
import pandas as pd
from typing import List, Tuple

# Load environment variables
load_dotenv()

# --- DIAGNOSTIC: Verify .env loading ---
print("\n--- .env Loading Diagnostics (Global Scope) ---")
dot_env_path = os.path.join(os.getcwd(), '.env')
if os.path.exists(dot_env_path):
    print(f".env file found at: {dot_env_path}")
    if load_dotenv(verbose=True, override=True):
        print("load_dotenv() reported successful loading/overriding.")
    else:
        print("load_dotenv() reported no new variables loaded or no override occurred.")
else:
    print(f".env file NOT found at: {dot_env_path}. This is likely the issue.")

LOCAL_ONEDRIVE_ROOT = os.getenv("LOCAL_ONEDRIVE_ROOT")

print(f"Value of LOCAL_ONEDRIVE_ROOT from os.getenv() (Global Scope): {repr(LOCAL_ONEDRIVE_ROOT)}")
print(f"Is LOCAL_ONEDRIVE_ROOT in os.environ (Global Scope)? {'LOCAL_ONEDRIVE_ROOT' in os.environ}")
print("--- End .env Loading Diagnostics (Global Scope) ---\n")


# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


# --- Webex API Configuration ---
WEBEX_API_URL = "https://webexapis.com/v1/messages"
AUTHORIZATION_TOKEN = os.getenv("WEBEX_AUTHORIZATION_TOKEN")
WEBEX_ROOM_ID = os.getenv("WEBEX_ROOM_ID")
FALLBACK_WEBEX_EMAIL = os.getenv("FALLBACK_EMAIL", "prcheruk@cisco.com")

if not AUTHORIZATION_TOKEN:
    logging.error("WEBEX_AUTHORIZATION_TOKEN environment variable not set. Webex communication may not work.")

# --- Oracle Database Connection Details ---
DB_USERNAME = os.getenv("DB_USERNAME")
DB_PASSWORD = os.getenv("DB_PASSWORD")
DB_DSN = (
    "(DESCRIPTION=(CONNECT_TIMEOUT=5)(TRANSPORT_CONNECT_TIMEOUT=3)(RETRY_COUNT=1)"
    "(ADDRESS_LIST=(LOAD_BALANCE=ON)(FAILOVER=ON)"
    "(ADDRESS=(PROTOCOL=TCP)(HOST=scan-prd-2102)(PORT=1541))"
    "(ADDRESS=(PROTOCOL=TCP)(HOST=scan-prd-2101)(PORT=1541)))"
    "(CONNECT_DATA=(SERVICE_NAME=CSFPRD_SRVC_RO.cisco.com)(SERVER=DEDICATED)))"
)

if not DB_USERNAME or not DB_PASSWORD or not DB_DSN:
    logging.error("DB_USERNAME, DB_PASSWORD, or DB_DSN environment variables not set. Exiting.")
    exit(1)

# --- Path to Oracle Instant Client ---
ORACLE_CLIENT_LIB_DIR = os.getenv("ORACLE_CLIENT_LIB_DIR")
try:
    if ORACLE_CLIENT_LIB_DIR and os.path.exists(ORACLE_CLIENT_LIB_DIR):
        try:
            cx_Oracle.init_oracle_client(lib_dir=ORACLE_CLIENT_LIB_DIR)
            logging.info(f"Initialized Oracle Client from: {ORACLE_CLIENT_LIB_DIR}")
        except cx_Oracle.Error as e:
            logging.warning(f"Failed to initialize Oracle Client from specified lib_dir '{ORACLE_CLIENT_LIB_DIR}': {e}. Attempting default initialization.")
            cx_Oracle.init_oracle_client()
            logging.info("Initialized Oracle Client using default search paths.")
    else:
        logging.warning("ORACLE_CLIENT_LIB_DIR not set or path invalid. Attempting default Oracle Client initialization.")
        cx_Oracle.init_oracle_client()
        logging.info("Initialized Oracle Client using default search paths.")
except cx_Oracle.Error as e:
    logging.error(f"FATAL: Error initializing Oracle Client (even with default search): {e}")
    logging.error("Please ensure Oracle Instant Client is correctly installed and configured, and its dependencies (like Visual C++ Redistributable) are met.")
    exit(1)

# --- Cisco AI Configuration ---
CISCO_AI_TOKEN_URL = "https://id.cisco.com/oauth2/default/v1/token"
CISCO_AI_CLIENT_ID = os.getenv("CISCO_AI_CLIENT_ID")
CISCO_AI_CLIENT_SECRET = os.getenv("CISCO_AI_CLIENT_SECRET")
CISCO_AI_APP_KEY = os.getenv("CISCO_AI_APP_KEY")
CISCO_AI_ENDPOINT = 'https://chat-ai.cisco.com'
CISCO_AI_API_VERSION = "2023-08-01-preview"
CISCO_AI_MODEL = "gpt-4.1" # Or other available model

if not all([CISCO_AI_CLIENT_ID, CISCO_AI_CLIENT_SECRET, CISCO_AI_APP_KEY]):
    logging.warning("One or more Cisco AI environment variables (CLIENT_ID, CLIENT_SECRET, APP_KEY) not set. AI summarization may not work.")


# --- Utility Functions ---

def json_serial(obj):
    """JSON serializer for objects not serializable by default json code"""
    if isinstance(obj, datetime):
        return obj.isoformat()
    raise TypeError(f"Type {type(obj)} not serializable}}")

def clean_excel_string(text):
    """
    Removes characters that are illegal in Excel worksheets (invalid XML characters).
    These typically include control characters.
    """
    if isinstance(text, str):
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    return text

def dicts_to_csv_string(data):
    """
    Converts a list of dictionaries to a CSV-formatted string.
    """
    if not data:
        return ""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=data[0].keys())
    writer.writeheader()
    writer.writerows(data)
    return output.getvalue()

def get_cisco_ai_client():
    """
    Obtains an OAuth token from Cisco and initializes the AzureOpenAI client.
    Returns the initialized client object or None on failure.
    """
    if not all([CISCO_AI_CLIENT_ID, CISCO_AI_CLIENT_SECRET, CISCO_AI_APP_KEY]):
        logging.error("Cisco AI credentials (CLIENT_ID, CLIENT_SECRET, APP_KEY) are incomplete. Cannot initialize AI client.")
        return None
    
    try:
        base64_auth = base64.b64encode(f'{CISCO_AI_CLIENT_ID}:{CISCO_AI_CLIENT_SECRET}'.encode('utf-8')).decode('utf-8')
        payload = "grant_type=client_credentials"
        headers = {
            "Accept": "*/*",
            "Content-Type": "application/x-www-form-urlencoded",
            "Authorization": f"Basic {base64_auth}"
        }

        logging.info("Attempting to get Cisco AI OAuth token...")
        token_response = requests.request("POST", CISCO_AI_TOKEN_URL, headers=headers, data=payload)
        token_response.raise_for_status()
        access_token = token_response.json()["access_token"]
        logging.info("Successfully obtained Cisco AI OAuth token.")

        client = openai.AzureOpenAI(
            azure_endpoint=CISCO_AI_ENDPOINT,
            api_key=access_token,
            api_version=CISCO_AI_API_VERSION
        )
        return client

    except requests.exceptions.RequestException as e:
        logging.error(f"Error obtaining Cisco AI token: {e}")
        if hasattr(e, 'response') and e.response is not None:
            logging.error(f"Response status code: {e.response.status_code}")
            logging.error(f"Response content: {e.response.text}")
        return None
    except Exception as e:
        logging.error(f"An unexpected error occurred during Cisco AI client initialization: {e}")
        return None

def summarize_with_cisco_ai(ai_client, text_to_summarize, prompt_prefix="Summarize the following data and provide actionable recommendations for a sales team:"):
    """
    Sends text to Cisco AI for summarization and recommendation.
    :param ai_client: An initialized openai.AzureOpenAI client.
    :param text_to_summarize: The analysis summary string to send to AI.
    :param prompt_prefix: An optional prefix for the AI prompt to guide summarization.
    :return: A string containing the AI-generated summary, or an error message.
    """
    if not ai_client:
        logging.warning("AI client is not initialized. Cannot summarize.")
        return "AI client is not initialized. Cannot summarize."

    if not text_to_summarize or not text_to_summarize.strip():
        logging.info("No data to summarize (text_to_summarize is empty).")
        return "No analysis data provided for AI summarization."

    messages = [
        {"role": "system", "content": "You are an expert Cisco Services Sales Consultant. Your task is to analyze customer data and provide concise, actionable recommendations for Services EA scope, highlighting opportunities and justifications."},
        {"role": "user", "content": f"{prompt_prefix}\n\n{text_to_summarize}"}
    ]

    try:
        logging.info("Sending data to Cisco AI for summarization...")
        response = ai_client.chat.completions.create(
            model=CISCO_AI_MODEL,
            messages=messages,
            user=f'{{"appkey": "{CISCO_AI_APP_KEY}"}}'
        )
        summary = response.choices[0].message.content
        
        if summary is None:
            logging.warning("Cisco AI returned None for summary content. This might mean the AI found nothing to summarize or an internal issue.")
            return "AI could not generate a summary for the provided data."
        elif not summary.strip():
            logging.warning("Cisco AI returned an empty or whitespace-only summary.")
            return "AI generated an empty summary. Data might be too sparse or prompt too restrictive."

        logging.info("Successfully received summary from Cisco AI.")
        return summary
    except openai.APIError as e:
        logging.error(f"Cisco AI API Error during summarization: {e}")
        if hasattr(e, 'response') and e.response is not None:
            logging.error(f"AI API Error Response status code: {e.response.status_code}")
            logging.error(f"AI API Error Response content: {e.response.text}")
        return f"Error summarizing data with AI: {e}"
    except Exception as e:
        logging.error(f"An unexpected error occurred during AI summarization: {e}")
        return f"Error summarizing data with AI: {e}"

def send_webex_message(to_email=None, room_id=None, text_message=None, card_payload=None):
    """
    Sends a Webex message, either plain text or an Adaptive Card, to a user and/or a room.
    """
    if not AUTHORIZATION_TOKEN:
        logging.error("WEBEX_AUTHORIZATION_TOKEN not set. Skipping sending Webex message.")
        return False

    if not AUTHORIZATION_TOKEN.startswith("Bearer "):
        bearer_token = f"Bearer {AUTHORIZATION_TOKEN}"
    else:
        bearer_token = AUTHORIZATION_TOKEN
    
    headers = {"Authorization": bearer_token}
    message_payload = {}

    target_identifier_for_log = None

    if room_id:
        message_payload["roomId"] = room_id
        target_identifier_for_log = f"Room ID: {repr(room_id)}"
    elif to_email:
        message_payload["toPersonEmail"] = to_email
        target_identifier_for_log = f"Email: {repr(to_email)}"
    else:
        logging.error("No recipient (to_email or room_id) specified for Webex message.")
        return False

    if card_payload:
        headers["Content-Type"] = "application/json"
        message_payload["text"] = text_message if text_message else "AI Analysis Summary"
        message_payload["attachments"] = [card_payload]
    elif text_message:
        headers["Content-Type"] = "application/json"
        message_payload["text"] = text_message
    else:
        logging.warning("No text message or card payload provided to send a Webex message.")
        return False

    print(f"\n--- Webex Send Diagnostic ---")
    print(f"Attempting to send to: {target_identifier_for_log}")
    print(f"Type of target ID: {type(room_id) if room_id else type(to_email)}")
    print(f"Length of target ID: {len(room_id) if room_id else len(to_email)}")
    print(f"Message Payload (first 500 chars): {json.dumps(message_payload, indent=2)[:500]}...")
    print(f"Headers: {headers}")
    print(f"--- End Webex Send Diagnostic ---\n")

    try:
        logging.info(f"Attempting to send Webex message to {room_id if room_id else to_email}...")
        response = requests.post(WEBEX_API_URL, headers=headers, data=json.dumps(message_payload))
        response.raise_for_status()
        logging.info(f"Successfully sent Webex message. Message ID: {response.json().get('id', 'N/A')}")
        return True
    except requests.exceptions.RequestException as e:
        logging.error(f"Error sending Webex message: {e}")
        if hasattr(e, 'response') and e.response is not None:
            logging.error(f"Webex API Response status code: {e.response.status_code}")
            logging.error(f"Webex API Response content: {e.response.text}")
        return False
    except Exception as e:
        logging.error(f"An unexpected error occurred during Webex message sending: {e}")
        return False


def execute_sql_file(sql_file_path, bind_params=None, db_username=DB_USERNAME, db_password=DB_PASSWORD, db_dsn=DB_DSN):
    """
    Executes an SQL query from a file, passing bind parameters.
    Returns a list of dictionaries, where each dictionary represents a row.
    """
    connection = None
    cursor = None
    results = []
    if bind_params is None:
        bind_params = {}

    try:
        with open(sql_file_path, 'r') as f:
            sql_content = f.read()
        logging.info(f"SQL query loaded from '{sql_file_path}'.")

        logging.info(f"Attempting to connect to Oracle database using DSN: {db_dsn}...")
        connection = cx_Oracle.connect(user=db_username, password=db_password, dsn=db_dsn)
        cursor = connection.cursor()
        logging.info("Successfully connected to Oracle database.")

        logging.info(f"Executing query with params: {bind_params}...")
        cursor.execute(sql_content, bind_params)
        logging.info(f"Query executed on DB server. Now fetching results for params: {bind_params}...")
        rows = cursor.fetchall()

        columns = [col[0].upper() for col in cursor.description]

        for row in rows:
            results.append(dict(zip(columns, row)))

        logging.info(f"Found {len(results)} records.")
        return results

    except FileNotFoundError:
        logging.error(f"Error: SQL file not found at '{sql_file_path}'.")
        return []
    except cx_Oracle.Error as e:
        error_obj, = e.args
        logging.error(f"Oracle Database Error ({error_obj.code}): {error_obj.message}")
        logging.error(f"Failed to execute query from '{sql_file_path}' with params '{bind_params}'.")
        return []
    except Exception as e:
        logging.error(f"An unexpected error occurred: {e}")
        return []
    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()
            logging.info("Database connection closed.")

def save_dataframe_to_csv(df, file_path, index=False):
    """
    Saves a pandas DataFrame to a CSV file.
    """
    try:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        df.to_csv(file_path, index=index)
        logging.info(f"DataFrame successfully saved to CSV: {file_path}")
        return True
    except Exception as e:
        logging.error(f"Error saving DataFrame to CSV '{file_path}': {e}")
        return False

def save_dataframe_to_excel(df, file_path, sheet_name="Data", index=False):
    """
    Saves a pandas DataFrame to an Excel file.
    """
    try:
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        df.to_excel(file_path, sheet_name=sheet_name, index=index)
        logging.info(f"DataFrame successfully saved to Excel: {file_path}")
        return True
    except Exception as e:
        logging.error(f"Error saving DataFrame to Excel '{file_path}': {e}")
        return False

def load_excel_from_local_onedrive(relative_onedrive_path, file_name, sheet_name=0):
    """
    Loads an Excel file from a locally synced OneDrive folder into a pandas DataFrame.
    """
    if not LOCAL_ONEDRIVE_ROOT:
        logging.error("LOCAL_ONEDRIVE_ROOT is not configured. Cannot load Excel from OneDrive.")
        return None
    if not os.path.isdir(LOCAL_ONEDRIVE_ROOT):
        logging.error(f"Configured LOCAL_ONEDRIVE_ROOT '{LOCAL_ONEDRIVE_ROOT}' is not a valid directory.")
        return None

    if relative_onedrive_path and (relative_onedrive_path.startswith('/') or relative_onedrive_path.startswith('\\')):
        logging.warning(f"Relative OneDrive path '{relative_onedrive_path}' starts with a separator. Removing it for correct joining.")
        relative_onedrive_path = relative_onedrive_path.lstrip('/\\')

    full_local_path = os.path.join(LOCAL_ONEDRIVE_ROOT, relative_onedrive_path, file_name)

    logging.info(f"Attempting to load Excel file from: {full_local_path}")

    if not os.path.exists(full_local_path):
        logging.error(f"Error: Excel file not found at '{full_local_path}'. "
                      f"Please ensure the file exists and your OneDrive is synced.")
        return None

    try:
        df = pd.read_excel(full_local_path, sheet_name=sheet_name)
        logging.info(f"Successfully loaded '{file_name}' into DataFrame (shape: {df.shape}).")
        return df
    except Exception as e:
        print(f"CRITICAL ERROR: Failed to read Excel file '{full_local_path}'. Exception: {e}")
        logging.error(f"Error loading Excel file '{file_name}' from '{full_local_path}': {e}")
        return None



def send_webex_message_with_files(to_email=None, room_id=None, text_message=None, file_paths=None, authorization_token=None):
    """
    Sends a Webex message with file attachments.
    """
    if not authorization_token:
        logging.error("WEBEX_AUTHORIZATION_TOKEN not set. Cannot send message.")
        return False
    
    if not authorization_token.startswith("Bearer "):
        bearer_token = f"Bearer {authorization_token}"
    else:
        bearer_token = authorization_token
    
    WEBEX_API_URL = "https://webexapis.com/v1/messages"
    
    # Prepare form data
    form_data = {}
    
    if room_id:
        form_data["roomId"] = room_id
    elif to_email:
        form_data["toPersonEmail"] = to_email
    else:
        logging.error("No recipient specified for Webex message.")
        return False
    
    if text_message:
        form_data["markdown"] = text_message
    
    # Prepare files
    files_to_upload = []
    if file_paths:
        for file_path in file_paths:
            try:
                file_obj = open(file_path, 'rb')
                file_name = file_path.split('/')[-1].split('\\')[-1]
                files_to_upload.append(('files', (file_name, file_obj, 'application/octet-stream')))
            except Exception as e:
                logging.error(f"Error opening file {file_path}: {e}")
    
    try:
        headers = {"Authorization": bearer_token}
        
        logging.info(f"Sending Webex message with {len(files_to_upload)} file(s)...")
        response = requests.post(
            WEBEX_API_URL,
            headers=headers,
            data=form_data,
            files=files_to_upload if files_to_upload else None
        )
        response.raise_for_status()
        
        logging.info(f"Successfully sent Webex message. Message ID: {response.json().get('id', 'N/A')}")
        
        # Close file handles
        for _, (_, file_obj, _) in files_to_upload:
            file_obj.close()
        
        return True
        
    except requests.exceptions.RequestException as e:
        logging.error(f"Error sending Webex message: {e}")
        if hasattr(e, 'response') and e.response is not None:
            logging.error(f"Response status: {e.response.status_code}")
            logging.error(f"Response content: {e.response.text}")
        
        for _, (_, file_obj, _) in files_to_upload:
            file_obj.close()
        
        return False


def prepare_data_for_openai(final_df: pd.DataFrame) -> str:
    """
    Convert the final analysis DataFrame into a structured text format for OpenAI.
    """
    summary_parts = []
    
    summary_parts.append("=== OVERALL STATISTICS ===")
    summary_parts.append(f"Total Parties Analyzed: {len(final_df)}")
    summary_parts.append(f"Total Install Base: {final_df['TOTAL_IB_COUNT'].sum():,.0f}")
    summary_parts.append(f"  - Covered: {final_df['IB_COUNT_COVERED'].sum():,.0f}")
    summary_parts.append(f"  - Uncovered: {final_df['IB_COUNT_UNCOVERED'].sum():,.0f}")
    summary_parts.append(f"  - Never Covered: {final_df['IB_COUNT_NEVER_COVERED'].sum():,.0f}")
    summary_parts.append(f"Total Support Cases: {final_df['CASE_COUNT'].sum():,.0f}")
    summary_parts.append(f"Total Sales Revenue: ${final_df['TOTAL_SALES'].sum():,.2f}")
    
    if 'EOL_PASSED' in final_df.columns:
        summary_parts.append("\n=== EOL/EOS RISK ANALYSIS ===")
        summary_parts.append(f"Assets Past EOL: {final_df['EOL_PASSED'].sum():,.0f}")
        summary_parts.append(f"Assets EOL Within 1 Year: {final_df['EOL_WITHIN_1YR'].sum():,.0f}")
        summary_parts.append(f"Assets Past EOS: {final_df['EOS_PASSED'].sum():,.0f}")
        summary_parts.append(f"Assets EOS Within 1 Year: {final_df['EOS_WITHIN_1YR'].sum():,.0f}")
    
    summary_parts.append("\n=== TOP 10 PARTIES BY INSTALL BASE ===")
    top_parties = final_df.nlargest(10, 'TOTAL_IB_COUNT')
    for idx, row in top_parties.iterrows():
        party_name = row.get('PARTY_NAME', f"Party {row['PARTY_ID']}")
        summary_parts.append(
            f"{party_name}: IB={row['TOTAL_IB_COUNT']:,.0f}, "
            f"Cases={row['CASE_COUNT']:,.0f}, Sales=${row['TOTAL_SALES']:,.2f}"
        )
    
    return "\n".join(summary_parts)

def create_ea_powerpoint_with_ai(
    final_df: pd.DataFrame,
    ai_summary: str,
    ai_client,
    cav_name: str,
    cav_id: int,
    gu_name: str,
    output_folder: str = "analysis_output"
) -> str:
    """
    Creates a PowerPoint presentation with AI-generated content and structure.
    """
    import os
    
    os.makedirs(output_folder, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_cav_name = cav_name.replace(" ", "_").replace("/", "_")
    filename = f"{safe_cav_name}_CAV_{cav_id}_EA_Presentation_{timestamp}.pptx"
    filepath = os.path.join(output_folder, filename)
    
    # Prepare data summary for AI
    data_summary = f"""
Customer: {gu_name}
CAV: {cav_name} (ID: {cav_id})

Key Metrics:
- Total Parties: {len(final_df)}
- Total Install Base: {final_df['TOTAL_IB_COUNT'].sum():,.0f}
  - Covered: {final_df['IB_COUNT_COVERED'].sum():,.0f}
  - Uncovered: {final_df['IB_COUNT_UNCOVERED'].sum():,.0f}
  - Never Covered: {final_df['IB_COUNT_NEVER_COVERED'].sum():,.0f}
- Total Cases: {final_df['CASE_COUNT'].sum():,.0f}
- Total Sales: ${final_df['TOTAL_SALES'].sum():,.2f}

EOL/EOS Risk:
- Assets Past EOL: {final_df.get('EOL_PASSED', pd.Series([0])).sum():,.0f}
- Assets EOL Within 1 Year: {final_df.get('EOL_WITHIN_1YR', pd.Series([0])).sum():,.0f}
- Assets Past EOS: {final_df.get('EOS_PASSED', pd.Series([0])).sum():,.0f}
- Assets EOS Within 1 Year: {final_df.get('EOS_WITHIN_1YR', pd.Series([0])).sum():,.0f}

Top 5 Customers by Install Base:
"""
    
    top_5 = final_df.nlargest(5, 'TOTAL_IB_COUNT')
    for idx, row in top_5.iterrows():
        party_name = row.get('PARTY_NAME', f"Party {row['PARTY_ID']}")
        data_summary += f"\n- {party_name}: {row['TOTAL_IB_COUNT']:,.0f} assets, {row['CASE_COUNT']:,.0f} cases"
    
    # Ask AI to generate PowerPoint content structure
    print("🤖 Asking AI to generate PowerPoint content...")
    
    ppt_prompt = f"""Based on this EA analysis data, create a PowerPoint presentation structure with 6-8 slides.

{data_summary}

Previous AI Analysis:
{ai_summary}

For each slide, provide:
1. Slide title
2. Key points (3-5 bullet points per slide)
3. Any specific recommendations or insights

Format your response as:
SLIDE 1: [Title]
- [Bullet point 1]
- [Bullet point 2]
...

SLIDE 2: [Title]
...

Focus on:
- Executive summary
- Key findings and metrics
- Risk analysis (EOL/EOS)
- Coverage opportunities
- Top customer insights
- Actionable recommendations
"""
    
    try:
        ppt_content_response = ai_client.chat.completions.create(
            model=CISCO_AI_MODEL,
            messages=[
                {"role": "system", "content": "You are an expert at creating executive presentations for Cisco Services sales. Create clear, actionable slide content."},
                {"role": "user", "content": ppt_prompt}
            ],
            user=f'{{"appkey": "{CISCO_AI_APP_KEY}"}}'
        )
        ppt_content = ppt_content_response.choices[0].message.content
        print("✅ AI generated PowerPoint content structure")
    except Exception as e:
        logging.error(f"Error getting AI PowerPoint content: {e}")
        ppt_content = f"SLIDE 1: Executive Summary\n{ai_summary}"
    
    # Parse AI response into slides
    slides_data = []
    current_slide = None
    
    for line in ppt_content.split('\n'):
        line = line.strip()
        if line.startswith('SLIDE'):
            if current_slide:
                slides_data.append(current_slide)
            # Extract title after "SLIDE X:"
            title = line.split(':', 1)[1].strip() if ':' in line else "Slide"
            current_slide = {'title': title, 'bullets': []}
        elif line.startswith('-') and current_slide:
            bullet = line.lstrip('- ').strip()
            if bullet:
                current_slide['bullets'].append(bullet)
    
    if current_slide:
        slides_data.append(current_slide)
    
    # Create PowerPoint
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = f"Enterprise Agreement Analysis"
    subtitle.text = f"{gu_name}\n{cav_name} (CAV ID: {cav_id})\n{datetime.now().strftime('%B %d, %Y')}"
    
    # Add AI-generated content slides
    for slide_data in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content layout
        
        title_shape = slide.shapes.title
        title_shape.text = slide_data['title']
        
        # Add content
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        for bullet in slide_data['bullets']:
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
            p.font.size = Pt(14)
            p.space_after = Pt(12)
    
    # Add a data visualization slide (Coverage Chart)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_shape.text_frame
    title_frame.text = "Install Base Coverage Distribution"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 141)
    
    chart_data = CategoryChartData()
    chart_data.categories = ['Covered', 'Uncovered', 'Never Covered']
    chart_data.add_series('IB Count', (
        final_df['IB_COUNT_COVERED'].sum(),
        final_df['IB_COUNT_UNCOVERED'].sum(),
        final_df['IB_COUNT_NEVER_COVERED'].sum()
    ))
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(1.5), Inches(1.5), Inches(7), Inches(5), chart_data
    ).chart
    chart.has_legend = True
    chart.legend.position = 2
    
    # Add metrics summary slide
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_shape.text_frame
    title_frame.text = "Key Metrics at a Glance"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 141)
    
    metrics = [
        ("Total Parties", f"{len(final_df):,}"),
        ("Total Install Base", f"{final_df['TOTAL_IB_COUNT'].sum():,.0f}"),
        ("IB Covered", f"{final_df['IB_COUNT_COVERED'].sum():,.0f}"),
        ("IB Uncovered", f"{final_df['IB_COUNT_UNCOVERED'].sum():,.0f}"),
        ("Total Cases", f"{final_df['CASE_COUNT'].sum():,.0f}"),
        ("Total Sales", f"${final_df['TOTAL_SALES'].sum():,.2f}")
    ]
    
    box_width, box_height = 2.8, 1.2
    start_x, start_y = 0.5, 1.2
    
    for i, (label, value) in enumerate(metrics):
        row, col = i // 3, i % 3
        x = start_x + col * (box_width + 0.3)
        y = start_y + row * (box_height + 0.3)
        
        box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(box_width), Inches(box_height))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(240, 240, 240)
        box.line.color.rgb = RGBColor(0, 51, 141)
        box.line.width = Pt(2)
        
        text_frame = box.text_frame
        text_frame.clear()
        
        p1 = text_frame.paragraphs[0]
        p1.text = label
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = text_frame.add_paragraph()
        p2.text = value
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(0, 51, 141)
        p2.alignment = PP_ALIGN.CENTER
    
    prs.save(filepath)
    print(f"✅ AI-powered PowerPoint presentation saved: {filepath}")
    return filepath

def create_ea_powerpoint(final_df: pd.DataFrame, ai_summary: str, cav_name: str, cav_id: int, gu_name: str, output_folder: str = "analysis_output") -> str:
    """
    Creates a PowerPoint presentation with EA analysis results and AI insights.
    """
    import os
    
    os.makedirs(output_folder, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_cav_name = cav_name.replace(" ", "_").replace("/", "_")
    filename = f"{safe_cav_name}_CAV_{cav_id}_EA_Presentation_{timestamp}.pptx"
    filepath = os.path.join(output_folder, filename)
    
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = f"Enterprise Agreement Analysis"
    subtitle.text = f"{gu_name}\n{cav_name} (CAV ID: {cav_id})\n{datetime.now().strftime('%B %d, %Y')}"
    
    # Slide 2: AI Summary
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_shape.text_frame
    title_frame.text = "Executive Summary"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 141)
    
    summary_shape = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(9), Inches(5.5))
    summary_frame = summary_shape.text_frame
    summary_frame.word_wrap = True
    summary_frame.text = ai_summary
    for paragraph in summary_frame.paragraphs:
        paragraph.font.size = Pt(12)
    
    # Slide 3: Key Metrics
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_shape.text_frame
    title_frame.text = "Key Metrics Overview"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 141)
    
    metrics = [
        ("Total Parties", f"{len(final_df):,}"),
        ("Total Install Base", f"{final_df['TOTAL_IB_COUNT'].sum():,.0f}"),
        ("IB Covered", f"{final_df['IB_COUNT_COVERED'].sum():,.0f}"),
        ("IB Uncovered", f"{final_df['IB_COUNT_UNCOVERED'].sum():,.0f}"),
        ("Total Cases", f"{final_df['CASE_COUNT'].sum():,.0f}"),
        ("Total Sales", f"${final_df['TOTAL_SALES'].sum():,.2f}")
    ]
    
    box_width, box_height = 2.8, 1.2
    start_x, start_y = 0.5, 1.2
    
    for i, (label, value) in enumerate(metrics):
        row, col = i // 3, i % 3
        x = start_x + col * (box_width + 0.3)
        y = start_y + row * (box_height + 0.3)
        
        box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(box_width), Inches(box_height))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(240, 240, 240)
        box.line.color.rgb = RGBColor(0, 51, 141)
        box.line.width = Pt(2)
        
        text_frame = box.text_frame
        text_frame.clear()
        
        p1 = text_frame.paragraphs[0]
        p1.text = label
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = text_frame.add_paragraph()
        p2.text = value
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(0, 51, 141)
        p2.alignment = PP_ALIGN.CENTER
    
    # Slide 4: Coverage Chart
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame = title_shape.text_frame
    title_frame.text = "Install Base Coverage Distribution"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(28)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 141)
    
    chart_data = CategoryChartData()
    chart_data.categories = ['Covered', 'Uncovered', 'Never Covered']
    chart_data.add_series('IB Count', (
        final_df['IB_COUNT_COVERED'].sum(),
        final_df['IB_COUNT_UNCOVERED'].sum(),
        final_df['IB_COUNT_NEVER_COVERED'].sum()
    ))
    
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(1.5), Inches(1.5), Inches(7), Inches(5), chart_data
    ).chart
    chart.has_legend = True
    chart.legend.position = 2
    
    prs.save(filepath)
    print(f"✅ PowerPoint presentation saved: {filepath}")
    return filepath

def perform_ea_analysis(
    target_cav_id_int: int,
    customer_detail_by_cav_df: pd.DataFrame,  # CHANGED: Now using CAV-specific customer data
    case_history_df: pd.DataFrame,
    install_base_df: pd.DataFrame,
    sales_history_df: pd.DataFrame,
    customer_serial_ib_df: pd.DataFrame,
    eamp_proposal_quote_df: pd.DataFrame,
) -> Tuple[pd.DataFrame, str]:
    """
    Perform EA analysis using CAV-specific customer data that matches IB/case/sales data.
    Returns (final_activity_summary_df, analysis_log_string).
    """
    messages: List[str] = []

    def _upper_cols(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        # Convert column names to strings first, then uppercase
        df.columns = [str(col).upper() for col in df.columns]
        return df

    def _safe_numeric(df: pd.DataFrame, cols: List[str]):
        for c in cols:
            if c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
        return df

    def _aggregate_ib_by_business_entity(ib_df: pd.DataFrame) -> pd.DataFrame:
        """Aggregate IB counts by U_BE (Business Entity) and U_SUBBE (Sub Business Entity)."""
        ib_df = _upper_cols(ib_df)
        
        # DEBUG: Print columns
        print(f"DEBUG: BE aggregation - columns after uppercase: {ib_df.columns.tolist()}")
        
        group_keys = ["CAV_BU_ID", "PARTY_ID"]
        
        # Check if required columns exist
        if not all(k in ib_df.columns for k in group_keys):
            missing_keys = [k for k in group_keys if k not in ib_df.columns]
            messages.append(f"Warning: BE aggregation skipped - missing keys {missing_keys}")
            return pd.DataFrame()
        
        if "U_BE" not in ib_df.columns and "U_SUBBE" not in ib_df.columns:
            messages.append("Warning: U_BE and U_SUBBE columns not found in install_base_df")
            return pd.DataFrame()
        
        # Convert merge keys to consistent types
        ib_df['CAV_BU_ID'] = pd.to_numeric(ib_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
        ib_df['PARTY_ID'] = pd.to_numeric(ib_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
        
        ib_df_valid = ib_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
        
        if ib_df_valid.empty:
            messages.append("Warning: No valid IB rows for BE aggregation")
            return pd.DataFrame()
        
        # Use QUANTITY if available
        if "QUANTITY" in ib_df_valid.columns:
            ib_df_valid["_CNT"] = pd.to_numeric(ib_df_valid["QUANTITY"], errors='coerce').fillna(1)
        else:
            ib_df_valid["_CNT"] = 1
        
        # Build breakdown strings per party
        result_data = []
        
        # Group by CAV_BU_ID and PARTY_ID to process each party
        for (cav_bu_id, party_id_val), party_data in ib_df_valid.groupby(group_keys):
            row = {"CAV_BU_ID": cav_bu_id, "PARTY_ID": party_id_val}
            
            # Build BE breakdown string
            be_breakdown = []
            if "U_BE" in party_data.columns:
                be_counts = party_data.groupby("U_BE")["_CNT"].sum().sort_values(ascending=False)
                be_breakdown = [f"{be}: {int(cnt)}" for be, cnt in be_counts.items() if pd.notna(be) and str(be).strip() != '']
            
            # Build Sub-BE breakdown string
            subbe_breakdown = []
            if "U_SUBBE" in party_data.columns:
                subbe_counts = party_data.groupby("U_SUBBE")["_CNT"].sum().sort_values(ascending=False)
                subbe_breakdown = [f"{subbe}: {int(cnt)}" for subbe, cnt in subbe_counts.items() if pd.notna(subbe) and str(subbe).strip() != '']
            
            row["BE_BREAKDOWN"] = "; ".join(be_breakdown) if be_breakdown else ""
            row["SUBBE_BREAKDOWN"] = "; ".join(subbe_breakdown) if subbe_breakdown else ""
            result_data.append(row)
        
        result = pd.DataFrame(result_data)
        
        if not result.empty:
            result['CAV_BU_ID'] = result['CAV_BU_ID'].astype(pd.Int64Dtype())
            result['PARTY_ID'] = result['PARTY_ID'].astype(pd.Int64Dtype())
            print(f"BE Summary: {len(result)} parties with BE/Sub-BE breakdowns")
            
            # Show sample breakdowns
            if "BE_BREAKDOWN" in result.columns:
                non_empty_be = result[result['BE_BREAKDOWN'] != '']
                if not non_empty_be.empty:
                    print(f"  Sample BE breakdown: {non_empty_be.iloc[0]['BE_BREAKDOWN'][:100]}...")
            if "SUBBE_BREAKDOWN" in result.columns:
                non_empty_subbe = result[result['SUBBE_BREAKDOWN'] != '']
                if not non_empty_subbe.empty:
                    print(f"  Sample Sub-BE breakdown: {non_empty_subbe.iloc[0]['SUBBE_BREAKDOWN'][:100]}...")
        
        return result

    def _aggregate_eol_eos_metrics(ib_df: pd.DataFrame) -> pd.DataFrame:
        """Aggregate EOL and EOS date metrics by party."""
        ib_df = _upper_cols(ib_df)
        
        # DEBUG: Print columns
        print(f"DEBUG: EOL/EOS aggregation - columns after uppercase: {ib_df.columns.tolist()}")
        
        group_keys = ["CAV_BU_ID", "PARTY_ID"]
        
        # Check if required columns exist
        if not all(k in ib_df.columns for k in group_keys):
            missing_keys = [k for k in group_keys if k not in ib_df.columns]
            messages.append(f"Warning: EOL/EOS aggregation skipped - missing keys {missing_keys}")
            return pd.DataFrame()
        
        if "EOL_DATE" not in ib_df.columns and "EOS_DATE" not in ib_df.columns:
            messages.append("Warning: EOL_DATE and EOS_DATE columns not found in install_base_df")
            return pd.DataFrame()
        
        # Convert merge keys to consistent types
        ib_df['CAV_BU_ID'] = pd.to_numeric(ib_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
        ib_df['PARTY_ID'] = pd.to_numeric(ib_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
        
        ib_df_valid = ib_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
        
        if ib_df_valid.empty:
            messages.append("Warning: No valid IB rows for EOL/EOS aggregation")
            return pd.DataFrame()
        
        # Convert date columns to datetime
        if "EOL_DATE" in ib_df_valid.columns:
            ib_df_valid["EOL_DATE"] = pd.to_datetime(ib_df_valid["EOL_DATE"], errors='coerce')
        if "EOS_DATE" in ib_df_valid.columns:
            ib_df_valid["EOS_DATE"] = pd.to_datetime(ib_df_valid["EOS_DATE"], errors='coerce')
        
        # Use QUANTITY if available
        if "QUANTITY" in ib_df_valid.columns:
            ib_df_valid["_CNT"] = pd.to_numeric(ib_df_valid["QUANTITY"], errors='coerce').fillna(1)
        else:
            ib_df_valid["_CNT"] = 1
        
        # Get current date for comparison
        current_date = pd.Timestamp.now()
        
        result_data = []
        
        # Group by CAV_BU_ID and PARTY_ID to process each party
        for (cav_bu_id, party_id_val), party_data in ib_df_valid.groupby(group_keys):
            row = {"CAV_BU_ID": cav_bu_id, "PARTY_ID": party_id_val}
            
            # EOL Metrics
            if "EOL_DATE" in party_data.columns:
                eol_data = party_data[party_data["EOL_DATE"].notna()]
                
                # Count assets by EOL status
                eol_passed = eol_data[eol_data["EOL_DATE"] < current_date]["_CNT"].sum()
                eol_upcoming_1yr = eol_data[(eol_data["EOL_DATE"] >= current_date) & 
                                            (eol_data["EOL_DATE"] <= current_date + pd.DateOffset(years=1))]["_CNT"].sum()
                eol_future = eol_data[eol_data["EOL_DATE"] > current_date + pd.DateOffset(years=1)]["_CNT"].sum()
                eol_unknown = party_data[party_data["EOL_DATE"].isna()]["_CNT"].sum()
                
                row["EOL_PASSED"] = int(eol_passed)
                row["EOL_WITHIN_1YR"] = int(eol_upcoming_1yr)
                row["EOL_FUTURE"] = int(eol_future)
                row["EOL_UNKNOWN"] = int(eol_unknown)
            else:
                row["EOL_PASSED"] = 0
                row["EOL_WITHIN_1YR"] = 0
                row["EOL_FUTURE"] = 0
                row["EOL_UNKNOWN"] = int(party_data["_CNT"].sum())
            
            # EOS Metrics
            if "EOS_DATE" in party_data.columns:
                eos_data = party_data[party_data["EOS_DATE"].notna()]
                
                # Count assets by EOS status
                eos_passed = eos_data[eos_data["EOS_DATE"] < current_date]["_CNT"].sum()
                eos_upcoming_1yr = eos_data[(eos_data["EOS_DATE"] >= current_date) & 
                                            (eos_data["EOS_DATE"] <= current_date + pd.DateOffset(years=1))]["_CNT"].sum()
                eos_future = eos_data[eos_data["EOS_DATE"] > current_date + pd.DateOffset(years=1)]["_CNT"].sum()
                eos_unknown = party_data[party_data["EOS_DATE"].isna()]["_CNT"].sum()
                
                row["EOS_PASSED"] = int(eos_passed)
                row["EOS_WITHIN_1YR"] = int(eos_upcoming_1yr)
                row["EOS_FUTURE"] = int(eos_future)
                row["EOS_UNKNOWN"] = int(eos_unknown)
            else:
                row["EOS_PASSED"] = 0
                row["EOS_WITHIN_1YR"] = 0
                row["EOS_FUTURE"] = 0
                row["EOS_UNKNOWN"] = int(party_data["_CNT"].sum())
            
            result_data.append(row)
        
        result = pd.DataFrame(result_data)
        
        if not result.empty:
            result['CAV_BU_ID'] = result['CAV_BU_ID'].astype(pd.Int64Dtype())
            result['PARTY_ID'] = result['PARTY_ID'].astype(pd.Int64Dtype())
            
            # Summary statistics
            total_eol_passed = result["EOL_PASSED"].sum()
            total_eol_upcoming = result["EOL_WITHIN_1YR"].sum()
            total_eos_passed = result["EOS_PASSED"].sum()
            total_eos_upcoming = result["EOS_WITHIN_1YR"].sum()
            
            print(f"EOL/EOS Summary: {len(result)} parties")
            print(f"  EOL: {total_eol_passed:,} passed | {total_eol_upcoming:,} within 1 year")
            print(f"  EOS: {total_eos_passed:,} passed | {total_eos_upcoming:,} within 1 year")
        
        return result

    def _aggregate_ib_from_install_base(ib_df: pd.DataFrame) -> pd.DataFrame:
        """Aggregate IB counts from install_base_df using QUANTITY and COVERAGE_STATUS."""
        # DEBUG: Check input dataframe
        print(f"DEBUG: install_base_df shape BEFORE uppercase: {ib_df.shape}")
        print(f"DEBUG: install_base_df columns BEFORE uppercase: {ib_df.columns.tolist()}")
        
        ib_df = _upper_cols(ib_df)
        
        # DEBUG: Print columns to diagnose issue
        print(f"DEBUG: install_base_df shape AFTER uppercase: {ib_df.shape}")
        print(f"DEBUG: install_base_df columns AFTER uppercase: {ib_df.columns.tolist()}")
        
        group_keys = ["CAV_BU_ID", "PARTY_ID"]
        ib_count_cols = ["IB_COUNT_COVERED", "IB_COUNT_UNCOVERED", "IB_COUNT_NEVER_COVERED"]

        if not all(k in ib_df.columns for k in group_keys):
            missing_keys = [k for k in group_keys if k not in ib_df.columns]
            messages.append(f"Warning: IB aggregation skipped - missing keys {missing_keys}. Available: {ib_df.columns.tolist()[:10]}")
            return pd.DataFrame(columns=group_keys + ib_count_cols)

        if "COVERAGE_STATUS" not in ib_df.columns:
            messages.append(f"Warning: IB aggregation skipped - missing COVERAGE_STATUS")
            return pd.DataFrame(columns=group_keys + ib_count_cols)

        # Convert merge keys to consistent types
        ib_df['CAV_BU_ID'] = pd.to_numeric(ib_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
        ib_df['PARTY_ID'] = pd.to_numeric(ib_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
        
        ib_df_valid = ib_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
        
        if ib_df_valid.empty:
            messages.append("Warning: No valid IB rows after removing null keys")
            return pd.DataFrame(columns=group_keys + ib_count_cols)

        def classify_coverage(val):
            if pd.isna(val):
                return "UNKNOWN"
            s = str(val).strip().upper()
            if s == "A":
                return "COVERED"
            if s == "I":
                return "UNCOVERED"
            if s == "N":
                return "NEVER_COVERED"
            return "UNKNOWN"

        ib_df_valid["_IB_CLASS"] = ib_df_valid["COVERAGE_STATUS"].apply(classify_coverage)
        
        # Use QUANTITY if available
        if "QUANTITY" in ib_df_valid.columns:
            ib_df_valid["_CNT"] = pd.to_numeric(ib_df_valid["QUANTITY"], errors='coerce').fillna(1)
            messages.append("DEBUG: Using QUANTITY column for IB counts")
        else:
            ib_df_valid["_CNT"] = 1
            messages.append("DEBUG: Using row count (1 per row) for IB counts")

        # Pivot to get counts per class
        pivot = ib_df_valid.pivot_table(index=group_keys, columns="_IB_CLASS", values="_CNT", aggfunc="sum", fill_value=0)
        pivot = pivot.reset_index()
        
        # FIX: Properly handle pivot result columns
        result = pivot.copy()
        
        # Ensure all classification columns exist with 0 defaults
        for col in ["COVERED", "UNCOVERED", "UNKNOWN", "NEVER_COVERED"]:
            if col not in result.columns:
                result[col] = 0
        
        # Create the final count columns
        result["IB_COUNT_COVERED"] = result["COVERED"]
        result["IB_COUNT_UNCOVERED"] = result["UNCOVERED"] + result["UNKNOWN"]
        result["IB_COUNT_NEVER_COVERED"] = result["NEVER_COVERED"]
        
        # Keep only the needed columns
        result = result[group_keys + ib_count_cols].copy()
        
        print(f"IB Aggregation: {len(result)} parties, Covered={result['IB_COUNT_COVERED'].sum():,.0f}, Uncovered={result['IB_COUNT_UNCOVERED'].sum():,.0f}, Never={result['IB_COUNT_NEVER_COVERED'].sum():,.0f}")
        
        return result

    def _aggregate_case_history(case_df: pd.DataFrame) -> pd.DataFrame:
        """Aggregate case counts from case_history_df."""
        case_df = _upper_cols(case_df)
        
        group_keys = ["CAV_BU_ID", "PARTY_ID"]
        
        if not all(k in case_df.columns for k in group_keys):
            messages.append(f"Warning: Case aggregation skipped - missing keys {group_keys}")
            return pd.DataFrame(columns=group_keys + ["CASE_COUNT"])
        
        case_df['CAV_BU_ID'] = pd.to_numeric(case_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
        case_df['PARTY_ID'] = pd.to_numeric(case_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
        
        case_df_valid = case_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
        
        if case_df_valid.empty:
            messages.append("Warning: No valid case rows")
            return pd.DataFrame(columns=group_keys + ["CASE_COUNT"])
        
        if "CASE_NUMBER" in case_df_valid.columns:
            result = case_df_valid.groupby(group_keys)["CASE_NUMBER"].nunique().reset_index()
            result.rename(columns={"CASE_NUMBER": "CASE_COUNT"}, inplace=True)
            messages.append("DEBUG: Using CASE_NUMBER for case counts (distinct count)")
        elif "CASE_ID" in case_df_valid.columns:
            result = case_df_valid.groupby(group_keys)["CASE_ID"].nunique().reset_index()
            result.rename(columns={"CASE_ID": "CASE_COUNT"}, inplace=True)
            messages.append("DEBUG: Using CASE_ID for case counts (distinct count)")
        else:
            result = case_df_valid.groupby(group_keys).size().reset_index(name="CASE_COUNT")
            messages.append("DEBUG: Using row count for case counts")
        
        print(f"Case Aggregation: {len(result)} parties, Total cases={result['CASE_COUNT'].sum():,.0f}")
        
        return result

    def _aggregate_sales_history(sales_df: pd.DataFrame) -> pd.DataFrame:
        """Aggregate sales amounts from sales_history_df."""
        sales_df = _upper_cols(sales_df)
        
        group_keys = ["CAV_BU_ID", "PARTY_ID"]
        
        if not all(k in sales_df.columns for k in group_keys):
            messages.append(f"Warning: Sales aggregation skipped - missing keys {group_keys}")
            return pd.DataFrame(columns=group_keys + ["TOTAL_SALES"])
        
        sales_df['CAV_BU_ID'] = pd.to_numeric(sales_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
        sales_df['PARTY_ID'] = pd.to_numeric(sales_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
        
        sales_df_valid = sales_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
        
        if sales_df_valid.empty:
            messages.append("Warning: No valid sales rows")
            return pd.DataFrame(columns=group_keys + ["TOTAL_SALES"])
        
        # Check for sales amount column - Added TOTAL_SALES_AMOUNT as first option
        amount_col = None
        for col_name in ["TOTAL_SALES_AMOUNT", "AMOUNT", "SALES_AMOUNT", "TOTAL_AMOUNT", "REVENUE", "SALES"]:
            if col_name in sales_df_valid.columns:
                amount_col = col_name
                break
        
        if amount_col is None:
            messages.append("Warning: No sales amount column found")
            result = sales_df_valid[group_keys].drop_duplicates()
            result["TOTAL_SALES"] = 0
            return result
        
        sales_df_valid[amount_col] = pd.to_numeric(sales_df_valid[amount_col], errors='coerce').fillna(0)
        result = sales_df_valid.groupby(group_keys)[amount_col].sum().reset_index()
        result.rename(columns={amount_col: "TOTAL_SALES"}, inplace=True)
        
        messages.append(f"DEBUG: Using {amount_col} column for sales amounts")
        print(f"Sales Aggregation: {len(result)} parties, Total sales=${result['TOTAL_SALES'].sum():,.2f}")
        
        return result

    # ===== MAIN PROCESSING =====
    
    # Step 1: Process CAV-specific customer base
    print("\n=== STEP 1: Processing CAV-specific customer base ===")
    base_df = _upper_cols(customer_detail_by_cav_df)
    
    if "CAV_BU_ID" not in base_df.columns or "PARTY_ID" not in base_df.columns:
        error_msg = f"Error: customer_detail_by_cav_df missing required merge keys CAV_BU_ID and/or PARTY_ID"
        print(error_msg)
        messages.append(error_msg)
        return pd.DataFrame(), "\n".join(messages)
    
    base_df['CAV_BU_ID'] = pd.to_numeric(base_df['CAV_BU_ID'], errors='coerce').astype(pd.Int64Dtype())
    base_df['PARTY_ID'] = pd.to_numeric(base_df['PARTY_ID'], errors='coerce').astype(pd.Int64Dtype())
    base_df = base_df.dropna(subset=['CAV_BU_ID', 'PARTY_ID']).copy()
    
    if base_df.empty:
        error_msg = "Error: No valid customers after removing null CAV_BU_ID/PARTY_ID"
        print(error_msg)
        messages.append(error_msg)
        return pd.DataFrame(), "\n".join(messages)
    
    # Keep customer info columns
    customer_cols = ["CAV_BU_ID", "PARTY_ID"]
    for col in ["CUSTOMER_NAME", "PARTY_NAME", "ACCOUNT_NAME", "CUSTOMER_ID", "PARTY_NUMBER"]:
        if col in base_df.columns:
            customer_cols.append(col)
    
    base_df = base_df[customer_cols].drop_duplicates(subset=["CAV_BU_ID", "PARTY_ID"])
    print(f"Customer base: {len(base_df)} unique party IDs for CAV {target_cav_id_int}")
    
    # Step 2-4: Aggregate data
    print("\n=== STEP 2-4: Aggregating IB, Cases, Sales ===")
    ib_agg = _aggregate_ib_from_install_base(install_base_df)
    be_agg = _aggregate_ib_by_business_entity(install_base_df)
    eol_eos_agg = _aggregate_eol_eos_metrics(install_base_df)
    case_agg = _aggregate_case_history(case_history_df)
    sales_agg = _aggregate_sales_history(sales_history_df)
    
    # Step 5: Check for PARTY_ID overlap
    print("\n=== STEP 5: PARTY_ID Overlap Check ===")
    base_party_ids = set(base_df['PARTY_ID'].dropna().astype(int))
    ib_party_ids = set(ib_agg['PARTY_ID'].dropna().astype(int)) if not ib_agg.empty else set()
    case_party_ids = set(case_agg['PARTY_ID'].dropna().astype(int)) if not case_agg.empty else set()
    sales_party_ids = set(sales_agg['PARTY_ID'].dropna().astype(int)) if not sales_agg.empty else set()
    
    print(f"Customer base: {len(base_party_ids)} party IDs")
    print(f"IB data: {len(ib_party_ids)} party IDs")
    print(f"Case data: {len(case_party_ids)} party IDs")
    print(f"Sales data: {len(sales_party_ids)} party IDs")
    print(f"Overlap (base ∩ IB): {len(base_party_ids & ib_party_ids)}")
    print(f"Overlap (base ∩ cases): {len(base_party_ids & case_party_ids)}")
    print(f"Overlap (base ∩ sales): {len(base_party_ids & sales_party_ids)}")
    
    # Step 6: Merge all aggregations
    print("\n=== STEP 6: Merging aggregations ===")
    merge_keys = ["CAV_BU_ID", "PARTY_ID"]
    
    final_df = base_df.copy()
    print(f"Starting with customer base: {len(final_df)} rows")
    
    # Merge IB
    if not ib_agg.empty:
        final_df = final_df.merge(ib_agg, on=merge_keys, how="left")
        for col in ["IB_COUNT_COVERED", "IB_COUNT_UNCOVERED", "IB_COUNT_NEVER_COVERED"]:
            if col in final_df.columns:
                final_df[col] = final_df[col].fillna(0)
        print(f"After IB merge: {len(final_df)} rows, Total IB = {final_df[['IB_COUNT_COVERED', 'IB_COUNT_UNCOVERED', 'IB_COUNT_NEVER_COVERED']].sum().sum():,.0f}")
    else:
        final_df["IB_COUNT_COVERED"] = 0
        final_df["IB_COUNT_UNCOVERED"] = 0
        final_df["IB_COUNT_NEVER_COVERED"] = 0
        messages.append("Warning: IB aggregation empty")
    
    # Merge cases
    if not case_agg.empty:
        final_df = final_df.merge(case_agg, on=merge_keys, how="left")
        if "CASE_COUNT" in final_df.columns:
            final_df["CASE_COUNT"] = final_df["CASE_COUNT"].fillna(0)
        print(f"After case merge: {len(final_df)} rows, Total cases = {final_df['CASE_COUNT'].sum():,.0f}")
    else:
        final_df["CASE_COUNT"] = 0
        messages.append("Warning: Case aggregation empty")
    
    # Merge sales
    if not sales_agg.empty:
        final_df = final_df.merge(sales_agg, on=merge_keys, how="left")
        if "TOTAL_SALES" in final_df.columns:
            final_df["TOTAL_SALES"] = final_df["TOTAL_SALES"].fillna(0)
        print(f"After sales merge: {len(final_df)} rows, Total sales = ${final_df['TOTAL_SALES'].sum():,.2f}")
    else:
        final_df["TOTAL_SALES"] = 0
        messages.append("Warning: Sales aggregation empty")
    
    # Merge BE/Sub-BE breakdowns
    if not be_agg.empty:
        final_df = final_df.merge(be_agg, on=merge_keys, how="left")
        if "BE_BREAKDOWN" in final_df.columns:
            final_df["BE_BREAKDOWN"] = final_df["BE_BREAKDOWN"].fillna("")
        if "SUBBE_BREAKDOWN" in final_df.columns:
            final_df["SUBBE_BREAKDOWN"] = final_df["SUBBE_BREAKDOWN"].fillna("")
        print(f"After BE merge: {len(final_df)} rows, Added BE/Sub-BE breakdowns")
    else:
        final_df["BE_BREAKDOWN"] = ""
        final_df["SUBBE_BREAKDOWN"] = ""
        messages.append("Warning: BE aggregation empty")
    
    # Merge EOL/EOS metrics
    if not eol_eos_agg.empty:
        final_df = final_df.merge(eol_eos_agg, on=merge_keys, how="left")
        eol_eos_cols = ["EOL_PASSED", "EOL_WITHIN_1YR", "EOL_FUTURE", "EOL_UNKNOWN",
                        "EOS_PASSED", "EOS_WITHIN_1YR", "EOS_FUTURE", "EOS_UNKNOWN"]
        for col in eol_eos_cols:
            if col in final_df.columns:
                final_df[col] = final_df[col].fillna(0)
        print(f"After EOL/EOS merge: {len(final_df)} rows, Added EOL/EOS metrics")
    else:
        eol_eos_cols = ["EOL_PASSED", "EOL_WITHIN_1YR", "EOL_FUTURE", "EOL_UNKNOWN",
                        "EOS_PASSED", "EOS_WITHIN_1YR", "EOS_FUTURE", "EOS_UNKNOWN"]
        for col in eol_eos_cols:
            final_df[col] = 0
        messages.append("Warning: EOL/EOS aggregation empty")
    
    # Step 7: Calculate totals
    print("\n=== STEP 7: Calculating totals ===")
    final_df["TOTAL_IB_COUNT"] = final_df["IB_COUNT_COVERED"] + final_df["IB_COUNT_UNCOVERED"] + final_df["IB_COUNT_NEVER_COVERED"]
    
    # Ensure numeric columns
    numeric_cols = ["IB_COUNT_COVERED", "IB_COUNT_UNCOVERED", "IB_COUNT_NEVER_COVERED", 
                    "TOTAL_IB_COUNT", "CASE_COUNT", "TOTAL_SALES",
                    "EOL_PASSED", "EOL_WITHIN_1YR", "EOL_FUTURE", "EOL_UNKNOWN",
                    "EOS_PASSED", "EOS_WITHIN_1YR", "EOS_FUTURE", "EOS_UNKNOWN"]
    final_df = _safe_numeric(final_df, numeric_cols)
    
    # Summary
    print(f"\n=== FINAL SUMMARY ===")
    print(f"Total parties: {len(final_df)}")
    print(f"Total IB Count: {final_df['TOTAL_IB_COUNT'].sum():,.0f}")
    print(f"  - Covered: {final_df['IB_COUNT_COVERED'].sum():,.0f}")
    print(f"  - Uncovered: {final_df['IB_COUNT_UNCOVERED'].sum():,.0f}")
    print(f"  - Never Covered: {final_df['IB_COUNT_NEVER_COVERED'].sum():,.0f}")
    print(f"Total Case Count: {final_df['CASE_COUNT'].sum():,.0f}")
    print(f"Total Sales: ${final_df['TOTAL_SALES'].sum():,.2f}")
    
    # Print BE/Sub-BE summary if available
    if "BE_BREAKDOWN" in final_df.columns and final_df["BE_BREAKDOWN"].str.len().sum() > 0:
        print(f"\nBusiness Entity breakdowns included for {(final_df['BE_BREAKDOWN'] != '').sum()} parties")
    if "SUBBE_BREAKDOWN" in final_df.columns and final_df["SUBBE_BREAKDOWN"].str.len().sum() > 0:
        print(f"Sub-Business Entity breakdowns included for {(final_df['SUBBE_BREAKDOWN'] != '').sum()} parties")
    
    # Print EOL/EOS summary if available
    if "EOL_PASSED" in final_df.columns:
        print(f"\nEOL Summary:")
        print(f"  - Passed EOL: {final_df['EOL_PASSED'].sum():,.0f} assets")
        print(f"  - EOL within 1 year: {final_df['EOL_WITHIN_1YR'].sum():,.0f} assets")
        print(f"  - EOL future (>1 year): {final_df['EOL_FUTURE'].sum():,.0f} assets")
        print(f"  - EOL unknown: {final_df['EOL_UNKNOWN'].sum():,.0f} assets")
    
    if "EOS_PASSED" in final_df.columns:
        print(f"\nEOS Summary:")
        print(f"  - Passed EOS: {final_df['EOS_PASSED'].sum():,.0f} assets")
        print(f"  - EOS within 1 year: {final_df['EOS_WITHIN_1YR'].sum():,.0f} assets")
        print(f"  - EOS future (>1 year): {final_df['EOS_FUTURE'].sum():,.0f} assets")
        print(f"  - EOS unknown: {final_df['EOS_UNKNOWN'].sum():,.0f} assets")
    
    messages.append(f"✓ Analysis complete: {len(final_df)} parties | Total IB: {final_df['TOTAL_IB_COUNT'].sum():,.0f} | Cases: {final_df['CASE_COUNT'].sum():,.0f} | Sales: ${final_df['TOTAL_SALES'].sum():,.2f}")
    
    log_output = "\n".join(messages)
    return final_df, log_output


# Separate optional function to save results to Excel
def save_ea_analysis_to_excel(
    final_df: pd.DataFrame,
    cav_name: str,
    cav_id: int,
    output_folder: str = "analysis_output"
) -> str:
    """
    Save the EA analysis results to a formatted Excel file.
    This is an OPTIONAL function - call it separately if you want Excel output.
    
    Args:
        final_df: The final analysis DataFrame returned from perform_ea_analysis()
        cav_name: Name of the CAV for filename
        cav_id: CAV ID for filename
        output_folder: Folder to save the output file
    
    Returns:
        str: Path to the saved file
    """
    import os
    from datetime import datetime
    
    # Create output folder if it doesn't exist
    os.makedirs(output_folder, exist_ok=True)
    
    # Generate filename with timestamp
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_cav_name = cav_name.replace(" ", "_").replace("/", "_")
    filename = f"{safe_cav_name}_CAV_{cav_id}_EA_Analysis_{timestamp}.xlsx"
    filepath = os.path.join(output_folder, filename)
    
    # Create Excel writer with xlsxwriter engine for formatting
    with pd.ExcelWriter(filepath, engine='xlsxwriter') as writer:
        # Write main data
        final_df.to_excel(writer, sheet_name='EA Analysis', index=False)
        
        # Get workbook and worksheet objects
        workbook = writer.book
        worksheet = writer.sheets['EA Analysis']
        
        # Define formats
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#4472C4',
            'font_color': 'white',
            'border': 1,
            'align': 'center',
            'valign': 'vcenter',
            'text_wrap': True
        })
        
        number_format = workbook.add_format({'num_format': '#,##0'})
        currency_format = workbook.add_format({'num_format': '$#,##0.00'})
        text_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})
        
        # Apply header format
        for col_num, col_name in enumerate(final_df.columns):
            worksheet.write(0, col_num, col_name, header_format)
        
        # Set column widths and formats
        col_formats = {
            'PARTY_ID': (15, number_format),
            'PARTY_NAME': (40, text_format),
            'CUSTOMER_NAME': (40, text_format),
            'IB_COUNT_COVERED': (18, number_format),
            'IB_COUNT_UNCOVERED': (18, number_format),
            'IB_COUNT_NEVER_COVERED': (22, number_format),
            'TOTAL_IB_COUNT': (18, number_format),
            'CASE_COUNT': (15, number_format),
            'TOTAL_SALES': (18, currency_format),
            'BE_BREAKDOWN': (60, text_format),
            'SUBBE_BREAKDOWN': (60, text_format),
            'EOL_PASSED': (15, number_format),
            'EOL_WITHIN_1YR': (18, number_format),
            'EOL_FUTURE': (15, number_format),
            'EOL_UNKNOWN': (15, number_format),
            'EOS_PASSED': (15, number_format),
            'EOS_WITHIN_1YR': (18, number_format),
            'EOS_FUTURE': (15, number_format),
            'EOS_UNKNOWN': (15, number_format),
        }
        
        for col_num, col_name in enumerate(final_df.columns):
            if col_name in col_formats:
                width, fmt = col_formats[col_name]
                worksheet.set_column(col_num, col_num, width, fmt)
            else:
                worksheet.set_column(col_num, col_num, 15)
        
        # Freeze first row (header)
        worksheet.freeze_panes(1, 0)
        
        # Add autofilter
        worksheet.autofilter(0, 0, len(final_df), len(final_df.columns) - 1)
        
        # Create summary sheet
        summary_data = {
            'Metric': [
                'Total Parties',
                'Total IB Count',
                '  - Covered',
                '  - Uncovered', 
                '  - Never Covered',
                'Total Case Count',
                'Total Sales',
                '',
                'EOL Passed',
                'EOL Within 1 Year',
                'EOL Future (>1 year)',
                'EOL Unknown',
                '',
                'EOS Passed',
                'EOS Within 1 Year',
                'EOS Future (>1 year)',
                'EOS Unknown'
            ],
            'Value': [
                len(final_df),
                final_df['TOTAL_IB_COUNT'].sum(),
                final_df['IB_COUNT_COVERED'].sum(),
                final_df['IB_COUNT_UNCOVERED'].sum(),
                final_df['IB_COUNT_NEVER_COVERED'].sum(),
                final_df['CASE_COUNT'].sum(),
                final_df['TOTAL_SALES'].sum(),
                '',
                final_df.get('EOL_PASSED', pd.Series([0])).sum(),
                final_df.get('EOL_WITHIN_1YR', pd.Series([0])).sum(),
                final_df.get('EOL_FUTURE', pd.Series([0])).sum(),
                final_df.get('EOL_UNKNOWN', pd.Series([0])).sum(),
                '',
                final_df.get('EOS_PASSED', pd.Series([0])).sum(),
                final_df.get('EOS_WITHIN_1YR', pd.Series([0])).sum(),
                final_df.get('EOS_FUTURE', pd.Series([0])).sum(),
                final_df.get('EOS_UNKNOWN', pd.Series([0])).sum()
            ]
        }
        
        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name='Summary', index=False)
        
        # Format summary sheet
        summary_ws = writer.sheets['Summary']
        summary_ws.set_column('A:A', 30, text_format)
        summary_ws.set_column('B:B', 20, number_format)
        
        # Apply header format to summary
        for col_num in range(2):
            summary_ws.write(0, col_num, summary_df.columns[col_num], header_format)
    
    print(f"✅ Excel file saved: {filepath}")
    return filepath
    
# --- Main execution logic ---
if __name__ == "__main__":
    # Define SQL file paths
    SQL_FILE_FOR_GU = "get_customers_by_gu_name.sql"
    # Commenting out SQL files for case history and install base
    # SQL_FILE_CASE_HISTORY = "get_case_history_by_cav.sql"
    # SQL_FILE_INSTALL_BASE = "get_install_base_by_cav.sql"

    # Define OneDrive paths and file names
    ONEDRIVE_SALES_HISTORY_FOLDER = "IB Analyst\\Metro Fire Dept" 
    SALES_HISTORY_EXCEL_FILE = "sales_history_data.xlsx" # Excel file name

    ONEDRIVE_PROPOSAL_INPUTS_FOLDER = "IB Analyst\\Metro Fire Dept"
    CUSTOMER_SERIAL_IB_EXCEL_FILE = "customer_serial_ib_data.xlsx"
    EAMP_PROPOSAL_QUOTE_EXCEL_FILE = "eamp_proposal_quote.xlsx"

    ONEDRIVE_CASE_INSTALL_BASE_FOLDER = "IB Analyst\\Metro Fire Dept"
    CASE_HISTORY_EXCEL_FILE = "case_history_by_cav.xlsx"
    INSTALL_BASE_EXCEL_FILE = "install_base_by_cav.xlsx"

    ONEDRIVE_CUSTOMER_DETAIL_FOLDER = "IB Analyst\\Metro Fire Dept"
    CUSTOMER_DETAIL_BY_CAV_EXCEL_FILE = "customer_detail_by_cavid.xlsx"


    # --- PROMPT TO CAPTURE GU_NAME ---
    target_gu_name = input("Enter the Global Ultimate Customer Name (e.g., 'Amazon'): ").strip()

    if not target_gu_name:
        logging.error("Global Ultimate Customer Name cannot be empty. Exiting.")
        exit(1)
    # --- END PROMPT TO CAPTURE GU_NAME ---

    # --- NEW: Define output_dir, sanitized_gu_name, timestamp early ---
    sanitized_gu_name = "".join(c for c in target_gu_name if c.isalnum() or c in (' ', '_')).replace(' ', '_')
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = "data_verification_output" # Folder to save files
    # --- END NEW ---

    print(f"\n--- Fetching Customer Records for Global Ultimate Customer Name: '{target_gu_name}' ---")
    customer_records_list = execute_sql_file(
        sql_file_path=SQL_FILE_FOR_GU,
        bind_params={'gu_name': target_gu_name}
    )

    if customer_records_list:
        print(f"Successfully retrieved {len(customer_records_list)} customer branch/subsidiary records for '{target_gu_name}'.")

        customer_df = pd.DataFrame(customer_records_list)
        customer_df.columns = customer_df.columns.str.upper()
        for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']:
            if col in customer_df.columns:
                customer_df[col] = pd.to_numeric(customer_df[col], errors='coerce').astype(pd.Int64Dtype())


        print("\n--- Customer Data Loaded into DataFrame ---")
        print("\nDataFrame Info:")
        customer_df.info()

        expected_cav_id_str = 'CAV_ID'
        expected_cav_name_str = 'CAV_NAME'
        is_cav_id_in_cols = expected_cav_id_str in customer_df.columns
        is_cav_name_in_cols = expected_cav_name_str in customer_df.columns

        if not is_cav_id_in_cols or not is_cav_name_in_cols:
            logging.error(f"CAV_ID or CAV_NAME column not found in DataFrame. Cannot aggregate by CAV. Please update SQL query '{SQL_FILE_FOR_GU}'.")
            exit(1)

        required_ib_cols = ['IB_COUNT_COVERED', 'IB_COUNT_UNCOVERED', 'IB_COUNT_NEVER_COVERED']
        for col in required_ib_cols:
            if col not in customer_df.columns:
                logging.warning(f"Column '{col}' not found in DataFrame. Please update SQL query '{SQL_FILE_FOR_GU}'. Setting to 0 for aggregation.")
                customer_df[col] = 0
            else:
                customer_df[col] = pd.to_numeric(customer_df[col], errors='coerce').fillna(0)

        cav_aggregated_df = customer_df.groupby(['CAV_ID', 'CAV_NAME'])[[
            'IB_COUNT_COVERED', 'IB_COUNT_UNCOVERED', 'IB_COUNT_NEVER_COVERED'
        ]].sum().reset_index()

        print(f"\n--- Aggregated Customer Data by CAV for '{target_gu_name}' ---")
        print(cav_aggregated_df.to_string())

        available_cav_ids = cav_aggregated_df['CAV_ID'].unique().tolist()
        if not available_cav_ids:
            logging.info("No CAV IDs found for further analysis.")
            exit(0)

        print(f"\nAvailable CAV IDs for further analysis: {', '.join(map(str, available_cav_ids))}")
        target_cav_id = input("Enter the CAV ID for detailed analysis: ").strip()

        if not target_cav_id:
            logging.error("CAV ID cannot be empty. Exiting.")
            exit(1)
        if str(target_cav_id) not in map(str, available_cav_ids):
            logging.error(f"Invalid CAV ID '{target_cav_id}'. Please select from available IDs.")
            exit(1)
        try:
            target_cav_id_int = int(target_cav_id)
        except ValueError:
            logging.error(f"CAV ID '{target_cav_id}' is not a valid number. Exiting.")
            exit(1)


        print(f"\n--- Performing Detailed Analysis for CAV ID: {target_cav_id} ---")

        # --- Load Case History from OneDrive Excel ---
        print("\nLoading Case History from OneDrive Excel...")
        case_history_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_CASE_INSTALL_BASE_FOLDER,
            file_name=CASE_HISTORY_EXCEL_FILE
        )
        if case_history_df is None: case_history_df = pd.DataFrame()
        else:
            case_history_df.columns = case_history_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']:
                if col in case_history_df.columns:
                    case_history_df[col] = pd.to_numeric(case_history_df[col], errors='coerce').astype(pd.Int64Dtype())
        if not case_history_df.empty and 'CAV_ID' in case_history_df.columns:
            case_history_df = case_history_df[case_history_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(case_history_df)} case history records (filtered for CAV ID {target_cav_id_int}).")


        # --- Load Install Base from OneDrive Excel ---
        print("\nLoading Install Base Data from OneDrive Excel...")
        install_base_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_CASE_INSTALL_BASE_FOLDER,
            file_name=INSTALL_BASE_EXCEL_FILE
        )
        if install_base_df is None: install_base_df = pd.DataFrame()
        else:
            install_base_df.columns = install_base_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']:
                if col in install_base_df.columns:
                    install_base_df[col] = pd.to_numeric(install_base_df[col], errors='coerce').astype(pd.Int64Dtype())
        if not install_base_df.empty and 'CAV_ID' in install_base_df.columns:
            install_base_df = install_base_df[install_base_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(install_base_df)} install base records (filtered for CAV ID {target_cav_id_int}).")


        # --- Read Sales History from Excel in OneDrive ---
        print("\nLoading Sales History from OneDrive Excel...")
        sales_history_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_SALES_HISTORY_FOLDER,
            file_name=SALES_HISTORY_EXCEL_FILE
        )
        if sales_history_df is None: sales_history_df = pd.DataFrame()
        else:
            sales_history_df.columns = sales_history_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']: # Assuming these might be in sales_history_df
                if col in sales_history_df.columns:
                    sales_history_df[col] = pd.to_numeric(sales_history_df[col], errors='coerce').astype(pd.Int64Dtype())
            # <<< END FIX >>>
        if not sales_history_df.empty and 'CAV_ID' in sales_history_df.columns:
            sales_history_df = sales_history_df[sales_history_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(sales_history_df)} sales history records (filtered for CAV ID {target_cav_id_int}).")


        # --- Load Customer-Provided Serial/Install Base Data from Excel ---
        print("\nLoading Customer-Provided Serial/Install Base Data from OneDrive Excel...")
        customer_serial_ib_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_PROPOSAL_INPUTS_FOLDER,
            file_name=CUSTOMER_SERIAL_IB_EXCEL_FILE
        )
        if customer_serial_ib_df is None: customer_serial_ib_df = pd.DataFrame()
        else:
            customer_serial_ib_df.columns = customer_serial_ib_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']: # Assuming these might be in customer_serial_ib_df
                if col in customer_serial_ib_df.columns:
                    customer_serial_ib_df[col] = pd.to_numeric(customer_serial_ib_df[col], errors='coerce').astype(pd.Int64Dtype())
            # <<< END FIX >>>
        if not customer_serial_ib_df.empty and 'CAV_ID' in customer_serial_ib_df.columns:
            customer_serial_ib_df = customer_serial_ib_df[customer_serial_ib_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(customer_serial_ib_df)} customer-provided serial/IB records (filtered for CAV ID {target_cav_id_int}).")


        # --- Load EAMP Proposal Quote Data from Excel ---
        print("\nLoading EAMP Proposal Quote Data from OneDrive Excel...")
        eamp_proposal_quote_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_PROPOSAL_INPUTS_FOLDER,
            file_name=EAMP_PROPOSAL_QUOTE_EXCEL_FILE
        )
        if eamp_proposal_quote_df is None: eamp_proposal_quote_df = pd.DataFrame()
        else:
            eamp_proposal_quote_df.columns = eamp_proposal_quote_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']: # Assuming these might be in eamp_proposal_quote_df
                if col in eamp_proposal_quote_df.columns:
                    eamp_proposal_quote_df[col] = pd.to_numeric(eamp_proposal_quote_df[col], errors='coerce').astype(pd.Int64Dtype())
            # <<< END FIX >>>
        if not eamp_proposal_quote_df.empty and 'CAV_ID' in eamp_proposal_quote_df.columns:
            eamp_proposal_quote_df = eamp_proposal_quote_df[eamp_proposal_quote_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(eamp_proposal_quote_df)} EAMP proposal quote records (filtered for CAV ID {target_cav_id_int}).")


        # --- Load Customer Hierarchy Data by CAV from Excel ---
        print("\nLoading Customer Hierarchy Data by CAV from OneDrive Excel...")
        customer_detail_by_cav_df = load_excel_from_local_onedrive(
            relative_onedrive_path=ONEDRIVE_CUSTOMER_DETAIL_FOLDER,
            file_name=CUSTOMER_DETAIL_BY_CAV_EXCEL_FILE
        )
        if customer_detail_by_cav_df is None: customer_detail_by_cav_df = pd.DataFrame()
        else:
            customer_detail_by_cav_df.columns = customer_detail_by_cav_df.columns.str.upper()
            for col in ['CAV_ID', 'CAV_BU_ID', 'PARTY_ID']:
                if col in customer_detail_by_cav_df.columns:
                    customer_detail_by_cav_df[col] = pd.to_numeric(customer_detail_by_cav_df[col], errors='coerce').astype(pd.Int64Dtype())
            # <<< END FIX >>>
        if not customer_detail_by_cav_df.empty and 'CAV_ID' in customer_detail_by_cav_df.columns:
            customer_detail_by_cav_df = customer_detail_by_cav_df[customer_detail_by_cav_df['CAV_ID'] == target_cav_id_int]
        print(f"Loaded {len(customer_detail_by_cav_df)} customer hierarchy records (filtered for CAV ID {target_cav_id_int}).")


        print("\n--- All requested data sets for detailed analysis are loaded. ---")
        
        # --- Perform Analysis ---
        final_summary_df_for_debug, ea_analysis_summary_str = perform_ea_analysis(
            target_cav_id_int,
            customer_detail_by_cav_df,
            case_history_df,
            install_base_df,
            sales_history_df,
            customer_serial_ib_df,
            eamp_proposal_quote_df
            )
        print("\n" + ea_analysis_summary_str) # Print the string part of the analysis

        # --- Save final_summary_df to Excel for verification ---
        if not final_summary_df_for_debug.empty:
            debug_summary_filename = f"{sanitized_gu_name}_CAV_{target_cav_id}_FinalSummary_DEBUG_{timestamp}.xlsx"
            debug_summary_filepath = os.path.join(output_dir, debug_summary_filename)
            save_dataframe_to_excel(final_summary_df_for_debug, debug_summary_filepath, sheet_name="Final Summary")
            print(f"\nDEBUG: Final summary DataFrame saved to: {debug_summary_filepath}")
        else:
            print("\nDEBUG: Final summary DataFrame is empty, not saving to Excel.")

        # --- Save final_summary_df to Excel for verification ---
        if not final_summary_df_for_debug.empty:
            debug_summary_filename = f"{sanitized_gu_name}_CAV_{target_cav_id}_FinalSummary_DEBUG_{timestamp}.xlsx"
            debug_summary_filepath = os.path.join(output_dir, debug_summary_filename)
            save_dataframe_to_excel(final_summary_df_for_debug, debug_summary_filepath, sheet_name="Final Summary")
            print(f"\nDEBUG: Final summary DataFrame saved to: {debug_summary_filepath}")
        else:
            print("\nDEBUG: Final summary DataFrame is empty, not saving to Excel.")

        # --- AI Summarization, Excel, PowerPoint, and Webex Delivery ---
        ai_input_text = prepare_data_for_openai(final_summary_df_for_debug)
        
        cisco_ai_client = get_cisco_ai_client()
        if cisco_ai_client:
            # ... rest of your code
      # --- AI Summarization, Excel, PowerPoint, and Webex Delivery ---
         ai_input_text = prepare_data_for_openai(final_summary_df_for_debug)

        cisco_ai_client = get_cisco_ai_client()
        if cisco_ai_client:
            logging.info("Sending analysis to AI...")
            ai_generated_summary = summarize_with_cisco_ai(
                ai_client=cisco_ai_client,
                text_to_summarize=ai_input_text,
                prompt_prefix=f"Analyze this EA data for {target_gu_name} (CAV: {target_cav_id_int}). Provide: 1) Executive summary, 2) Coverage opportunities, 3) EOL/EOS recommendations, 4) Top customers, 5) Next steps."
            )
            print("\n--- AI Summary ---")
            print(ai_generated_summary)
            
            # Generate Excel
            excel_path = save_ea_analysis_to_excel(
                final_df=final_summary_df_for_debug,
                cav_name=target_gu_name,
                cav_id=target_cav_id_int,
                output_folder=output_dir
            )
            
            # Generate PowerPoint
            #ppt_path = create_ea_powerpoint(
            #    final_df=final_summary_df_for_debug,
            #    ai_summary=ai_generated_summary,
            #    cav_name=target_gu_name,
            #    cav_id=target_cav_id_int,
            #    gu_name=target_gu_name,
            #    output_folder=output_dir
            #)

            # Generate PowerPoint with AI-generated content
            ppt_path = create_ea_powerpoint_with_ai(
                final_df=final_summary_df_for_debug,
                ai_summary=ai_generated_summary,
                ai_client=cisco_ai_client,  # Pass the AI client
                cav_name=target_gu_name,
                cav_id=target_cav_id_int,
                gu_name=target_gu_name,
                output_folder=output_dir
            )
            
            # Send to Webex
            webex_message = f"""
# ✅ EA Analysis Complete: {target_gu_name}

## 📊 Key Metrics
- **Parties:** {len(final_summary_df_for_debug):,}
- **Install Base:** {final_summary_df_for_debug['TOTAL_IB_COUNT'].sum():,.0f}
- **Cases:** {final_summary_df_for_debug['CASE_COUNT'].sum():,.0f}
- **Sales:** ${final_summary_df_for_debug['TOTAL_SALES'].sum():,.2f}

## 🤖 AI Insights
{ai_generated_summary}

## 📁 Files Attached
1. Excel Report
2. PowerPoint Presentation
            """
            
            if AUTHORIZATION_TOKEN:
                # Send message with AI summary and Excel file
                send_webex_message_with_files(
                    room_id=WEBEX_ROOM_ID,
                    to_email=FALLBACK_WEBEX_EMAIL,
                    text_message=webex_message.strip(),
                    file_paths=[excel_path],
                    authorization_token=AUTHORIZATION_TOKEN
                )
                
                # Send PowerPoint in a separate message
                send_webex_message_with_files(
                    room_id=WEBEX_ROOM_ID,
                    to_email=FALLBACK_WEBEX_EMAIL,
                    text_message="📊 PowerPoint Presentation",
                    file_paths=[ppt_path],
                    authorization_token=AUTHORIZATION_TOKEN
                )
                print(f"\n✅ Sent to Webex: AI Summary + Excel + PowerPoint")
        else:
            print("⚠️ AI client failed. Generating files only.")
            excel_path = save_ea_analysis_to_excel(final_summary_df_for_debug, target_gu_name, target_cav_id_int, output_dir)
            ppt_path = create_ea_powerpoint(final_summary_df_for_debug, "AI unavailable", target_gu_name, target_cav_id_int, target_gu_name, output_dir)

    else:
        print(f"No customer records found for '{target_gu_name}'.")

    print("\n--- Script finished ---")